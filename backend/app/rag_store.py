from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Awaitable, Callable, Iterable, Sequence

import numpy as np


@dataclass(frozen=True)
class IngestReport:
    ingested: int
    updated: int
    skipped: int
    failed: int
    errors: list[str]


@dataclass(frozen=True)
class RagDocument:
    doc_id: str
    file_path: str
    title: str
    chunk_count: int
    updated_at: float


@dataclass(frozen=True)
class RetrievedChunk:
    doc_id: str
    file_path: str
    title: str
    chunk_index: int
    score: float
    text: str


class LocalRagStore:
    def __init__(
        self,
        db_path: str,
        chunk_size_chars: int = 1200,
        chunk_overlap_chars: int = 180,
    ) -> None:
        self.db_path = Path(db_path)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.chunk_size_chars = max(300, int(chunk_size_chars))
        self.chunk_overlap_chars = max(0, min(int(chunk_overlap_chars), self.chunk_size_chars // 2))
        self._conn = sqlite3.connect(self.db_path)
        self._conn.row_factory = sqlite3.Row
        self._init_schema()

    def close(self) -> None:
        self._conn.close()

    def _init_schema(self) -> None:
        self._conn.executescript(
            """
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS documents (
                doc_id TEXT PRIMARY KEY,
                file_path TEXT NOT NULL UNIQUE,
                title TEXT NOT NULL,
                sha256 TEXT NOT NULL,
                file_size INTEGER NOT NULL,
                file_mtime REAL NOT NULL,
                chunk_count INTEGER NOT NULL,
                updated_at REAL NOT NULL
            );

            CREATE TABLE IF NOT EXISTS chunks (
                chunk_id TEXT PRIMARY KEY,
                doc_id TEXT NOT NULL,
                chunk_index INTEGER NOT NULL,
                text TEXT NOT NULL,
                vector BLOB NOT NULL,
                dim INTEGER NOT NULL,
                norm REAL NOT NULL,
                FOREIGN KEY(doc_id) REFERENCES documents(doc_id) ON DELETE CASCADE
            );

            CREATE INDEX IF NOT EXISTS idx_chunks_doc_id ON chunks(doc_id);
            CREATE INDEX IF NOT EXISTS idx_documents_updated_at ON documents(updated_at DESC);
            """
        )
        self._conn.commit()

    async def ingest_paths(
        self,
        file_paths: Sequence[str],
        embed_texts: Callable[[Sequence[str]], Awaitable[list[list[float]]]],
    ) -> IngestReport:
        normalized = _normalize_paths(file_paths)
        errors: list[str] = []
        ingested = 0
        updated = 0
        skipped = 0
        failed = 0

        for file_path in normalized:
            try:
                state = os.stat(file_path)
            except OSError as exc:
                failed += 1
                errors.append(f"{file_path}: {exc}")
                continue

            try:
                text = _extract_text(Path(file_path))
                if not text.strip():
                    failed += 1
                    errors.append(f"{file_path}: no extractable text")
                    continue
                text_hash = hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()
                existing = self._conn.execute(
                    "SELECT doc_id, sha256 FROM documents WHERE file_path = ?",
                    (file_path,),
                ).fetchone()
                if existing and str(existing["sha256"]) == text_hash:
                    skipped += 1
                    continue

                chunks = _chunk_text(text, self.chunk_size_chars, self.chunk_overlap_chars)
                if not chunks:
                    failed += 1
                    errors.append(f"{file_path}: no chunks generated")
                    continue

                vectors = await embed_texts(chunks)
                if len(vectors) != len(chunks):
                    failed += 1
                    errors.append(f"{file_path}: embedding count mismatch")
                    continue

                doc_id = hashlib.sha1(file_path.encode("utf-8")).hexdigest()
                now_ts = time.time()
                if existing:
                    self._conn.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
                self._conn.execute(
                    """
                    INSERT INTO documents(doc_id, file_path, title, sha256, file_size, file_mtime, chunk_count, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(doc_id) DO UPDATE SET
                        file_path=excluded.file_path,
                        title=excluded.title,
                        sha256=excluded.sha256,
                        file_size=excluded.file_size,
                        file_mtime=excluded.file_mtime,
                        chunk_count=excluded.chunk_count,
                        updated_at=excluded.updated_at
                    """,
                    (
                        doc_id,
                        file_path,
                        Path(file_path).name,
                        text_hash,
                        int(state.st_size),
                        float(state.st_mtime),
                        len(chunks),
                        now_ts,
                    ),
                )

                for chunk_index, (chunk_text, vector) in enumerate(zip(chunks, vectors)):
                    arr = np.asarray(vector, dtype=np.float32)
                    norm = float(np.linalg.norm(arr))
                    if not np.isfinite(norm) or norm <= 0.0:
                        continue
                    chunk_id = f"{doc_id}:{chunk_index}"
                    self._conn.execute(
                        """
                        INSERT OR REPLACE INTO chunks(chunk_id, doc_id, chunk_index, text, vector, dim, norm)
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            chunk_id,
                            doc_id,
                            chunk_index,
                            chunk_text,
                            arr.tobytes(),
                            int(arr.shape[0]),
                            norm,
                        ),
                    )

                self._conn.commit()
                if existing:
                    updated += 1
                else:
                    ingested += 1
            except Exception as exc:  # pragma: no cover - defensive ingestion guard
                failed += 1
                errors.append(f"{file_path}: {exc}")

        return IngestReport(
            ingested=ingested,
            updated=updated,
            skipped=skipped,
            failed=failed,
            errors=errors,
        )

    def list_documents(self) -> list[RagDocument]:
        rows = self._conn.execute(
            """
            SELECT doc_id, file_path, title, chunk_count, updated_at
            FROM documents
            ORDER BY updated_at DESC
            """
        ).fetchall()
        return [
            RagDocument(
                doc_id=str(row["doc_id"]),
                file_path=str(row["file_path"]),
                title=str(row["title"]),
                chunk_count=int(row["chunk_count"]),
                updated_at=float(row["updated_at"]),
            )
            for row in rows
        ]

    def clear(self) -> None:
        self._conn.execute("DELETE FROM chunks")
        self._conn.execute("DELETE FROM documents")
        self._conn.commit()

    async def search(
        self,
        query: str,
        top_k: int,
        embed_texts: Callable[[Sequence[str]], Awaitable[list[list[float]]]],
    ) -> list[RetrievedChunk]:
        query = (query or "").strip()
        if not query:
            return []
        embedded = await embed_texts([query])
        if not embedded:
            return []
        query_vec = np.asarray(embedded[0], dtype=np.float32)
        query_norm = float(np.linalg.norm(query_vec))
        if not np.isfinite(query_norm) or query_norm <= 0.0:
            return []

        rows = self._conn.execute(
            """
            SELECT
                c.doc_id,
                c.chunk_index,
                c.text,
                c.vector,
                c.dim,
                c.norm,
                d.file_path,
                d.title
            FROM chunks c
            JOIN documents d ON d.doc_id = c.doc_id
            """
        ).fetchall()

        scored: list[RetrievedChunk] = []
        for row in rows:
            dim = int(row["dim"])
            if dim != int(query_vec.shape[0]):
                continue
            norm = float(row["norm"])
            if not np.isfinite(norm) or norm <= 0.0:
                continue
            vec = np.frombuffer(row["vector"], dtype=np.float32)
            if vec.shape[0] != dim:
                continue
            score = float(np.dot(query_vec, vec) / (query_norm * norm))
            if not np.isfinite(score):
                continue
            scored.append(
                RetrievedChunk(
                    doc_id=str(row["doc_id"]),
                    file_path=str(row["file_path"]),
                    title=str(row["title"]),
                    chunk_index=int(row["chunk_index"]),
                    score=score,
                    text=str(row["text"]),
                )
            )

        scored.sort(key=lambda item: item.score, reverse=True)
        return scored[: max(1, int(top_k))]


def _normalize_paths(raw_paths: Sequence[str]) -> list[str]:
    unique: list[str] = []
    seen: set[str] = set()
    for raw in raw_paths:
        candidate = str(raw or "").strip()
        if not candidate:
            continue
        resolved = str(Path(candidate).expanduser().resolve())
        if resolved in seen:
            continue
        path_obj = Path(resolved)
        if path_obj.is_dir():
            for suffix in ("*.txt", "*.md", "*.pdf", "*.docx"):
                for child in sorted(path_obj.rglob(suffix)):
                    child_resolved = str(child.resolve())
                    if child_resolved not in seen:
                        seen.add(child_resolved)
                        unique.append(child_resolved)
            continue
        seen.add(resolved)
        unique.append(resolved)
    return unique


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".rst", ".log", ".json"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages: list[str] = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(page_text.strip())
        return "\n\n".join(pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        blocks = [para.text.strip() for para in document.paragraphs if para.text and para.text.strip()]
        return "\n".join(blocks)
    raise ValueError(f"unsupported file extension: {suffix or '<none>'}")


def _chunk_text(text: str, chunk_size_chars: int, chunk_overlap_chars: int) -> list[str]:
    clean = "\n".join(line.rstrip() for line in text.splitlines())
    clean = "\n".join(line for line in clean.splitlines() if line.strip())
    if not clean:
        return []
    paragraphs = [p.strip() for p in clean.split("\n\n") if p.strip()]
    if not paragraphs:
        return []

    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= chunk_size_chars:
            current = candidate
            continue
        if current:
            chunks.append(current)
        if len(paragraph) <= chunk_size_chars:
            current = paragraph
            continue
        start = 0
        stride = max(1, chunk_size_chars - chunk_overlap_chars)
        while start < len(paragraph):
            part = paragraph[start : start + chunk_size_chars].strip()
            if part:
                chunks.append(part)
            start += stride
        current = ""

    if current:
        chunks.append(current)
    return chunks
